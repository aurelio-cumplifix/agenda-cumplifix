# -*- coding: utf-8 -*-
"""
Contrasta el dataset generado contra las fechas declaradas literalmente en el PPTX fuente.
Cualquier discrepancia se reporta como OBSERVACIÓN para validación humana.
"""
import json, re, unicodedata
from pptx import Presentation

PPTX = "/root/.claude/uploads/a114e646-8ca5-5b6f-b22a-266224d9c687/ea3caae5-AGENDA_DE_CUMPLIMIENTO_ANUAL_CUMPLIFIX_CONDUSEF_2026.pptx"
data = json.load(open("agenda-2026.json", encoding="utf-8"))
MESES = ["enero","febrero","marzo","abril","mayo","junio","julio","agosto",
         "septiembre","octubre","noviembre","diciembre"]

def norm(s):
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode().lower()
    return re.sub(r"\s+", " ", s).strip()

# --- 1. Recolectar de las tablas del PPTX los rangos "Del X al Y de MES" y "Solo el X"
prs = Presentation(PPTX)
declarado = {}   # mes -> set de textos de fecha límite
for s in prs.slides:
    for sh in s.shapes:
        if not sh.has_table:
            continue
        t = sh.table
        enc = [norm(c.text) for c in t.rows[0].cells]
        if "fecha limite" not in enc:
            continue
        i = enc.index("fecha limite")
        j = enc.index("periodicidad") if "periodicidad" in enc else None
        for r in list(t.rows)[1:]:
            txt = norm(r.cells[i].text)
            if not txt:
                continue
            m = re.search(r"de ([a-z]+) de 2026", txt) or re.search(r"al \d+ de ([a-z]+)", txt)
            if not m or m.group(1) not in MESES:
                continue
            declarado.setdefault(MESES.index(m.group(1)) + 1, set()).add(txt)

obs = []
print("=" * 78)
print("VERIFICACIÓN — dataset generado vs. PPTX fuente")
print("=" * 78)

for mes in sorted(declarado):
    md = next(x for x in data["meses"] if x["mes"] == mes)
    gen = set(norm(o["limite"]) for o in md["obligaciones"])
    # Normalizamos "unicamente el X de mes de 2026" -> "solo el X de mes de 2026"
    def key(t):
        t = t.replace("unicamente el", "solo el").replace("solo el", "solo el")
        return re.sub(r"[^a-z0-9 ]", "", t)
    gen_k = set(key(g) for g in gen)
    ok, dif = [], []
    for d in sorted(declarado[mes]):
        dk = key(d)
        if dk in gen_k:
            ok.append(d)
        else:
            dif.append(d)
    estado = "OK" if not dif else "REVISAR"
    print("\n{:<11} [{}]  coinciden {}/{}".format(MESES[mes-1], estado, len(ok), len(ok)+len(dif)))
    for d in dif:
        print("   · PPTX dice : {}".format(d))
        print("     generado  : {}".format(sorted(gen_k)))
        obs.append((MESES[mes-1], d))

# --- 2. Coherencia interna de Art. 160 / Art. 161 en el PPTX
print("\n" + "=" * 78)
print("COHERENCIA DEL CATÁLOGO EN LA FUENTE (Art. 160 vs Art. 161)")
print("=" * 78)
for s in prs.slides:
    for sh in s.shapes:
        if not sh.has_table:
            continue
        t = sh.table
        enc = [norm(c.text) for c in t.rows[0].cells]
        if "fecha limite" not in enc:
            continue
        cols = {n: k for k, n in enumerate(enc)}
        ic = cols.get("obligacion que reportar de la d.m.r.c.", cols.get("fundamento legal"))
        ip = cols.get("periodo que se reporta") or cols.get("periodo que se reporta 🪅")
        if ic is None or ip is None:
            continue
        for r in list(t.rows)[1:]:
            art = norm(r.cells[ic].text)
            per = norm(r.cells[ip].text)
            lim = norm(r.cells[cols["fecha limite"]].text)
            m = re.search(r"al \d+ de ([a-z]+)|de ([a-z]+) de 2026", lim)
            if not m:
                continue
            mesnom = m.group(1) or m.group(2)
            if mesnom not in MESES:
                continue
            if art.startswith("art. 160"):
                esperado = "mes de " + mesnom + " de 2026"
                if per != esperado:
                    print("  Art. 160 en {:<10} → periodo declarado '{}' (se esperaría '{}')".format(mesnom, per, esperado))
            if art.startswith("art. 161"):
                if per == "mes de " + mesnom + " de 2026":
                    print("  Art. 161 en {:<10} → periodo declarado '{}' (se esperaría el mes anterior)".format(mesnom, per))

print("\n" + "=" * 78)
print("RESUMEN: {} discrepancia(s) de fecha límite".format(len(obs)))
